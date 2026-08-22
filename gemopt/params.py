"""The search space, and the spreadsheet that defines it.

The client already works through Excel, so the list of parameters, their
allowed ranges and the rules that tie them together live in a spreadsheet he
can edit without touching any code.  Everything the optimiser is allowed to
move is in that one file; nothing is hidden in the source.

Reading and writing both work with plain CSV as well, so the tooling does not
hard-depend on openpyxl.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from .geometry import CutParams

__all__ = ["ParamSpec", "SearchSpace", "DEFAULT_SPACE", "write_template", "load_space"]


@dataclass
class ParamSpec:
    """One row of the spreadsheet."""

    name: str
    low: float
    high: float
    default: float
    kind: str = "float"        # float | int | bool
    step: float = 0.0          # 0 = continuous
    enabled: bool = True
    note: str = ""

    @property
    def fixed(self) -> bool:
        return (not self.enabled) or self.low >= self.high

    def clamp(self, value: float) -> float:
        v = float(np.clip(value, self.low, self.high))
        if self.kind == "int":
            return float(int(round(v)))
        if self.kind == "bool":
            return float(v >= 0.5)
        if self.step > 0.0:
            return float(self.low + round((v - self.low) / self.step) * self.step)
        return v


# The ranges below are the proportions the trade considers cuttable for a
# round brilliant.  They are deliberately wider than the "ideal" window --
# the whole point of the project is that the ideal window for glass is not
# the same as the one for diamond, so the search has to be allowed outside
# it -- but not so wide that it wanders into shapes no cutter could make.
DEFAULT_SPACE: list[ParamSpec] = [
    ParamSpec("n_base", 6, 12, 8, kind="int", step=1, enabled=False,
              note="base symmetry; every row is a multiple of this. Off by default: change it with --facets"),
    ParamSpec("girdle_multiple", 1, 4, 2, kind="int", step=1, enabled=False,
              note="girdle facets per main facet. Off by default: change it with --girdle-multiple"),
    ParamSpec("table", 0.45, 0.68, 0.53,
              note="table radius / girdle radius (0.53 = 53% table)"),
    ParamSpec("crown_angle", 25.0, 42.0, 34.5,
              note="crown (bezel) angle in degrees"),
    # The usual trade range stops around 46 degrees, but that range was drawn
    # up for diamond.  A low-index material has a far wider critical angle and
    # its best pavilion sits much deeper -- capping at 46 quietly hides the
    # answer for glass, so the ceiling is left where the geometry, not the
    # diamond convention, runs out.
    ParamSpec("pavilion_angle", 36.0, 54.0, 40.75,
              note="pavilion main angle in degrees - the single most important number"),
    ParamSpec("star_length", 0.35, 0.75, 0.55,
              note="star facet length as a fraction"),
    ParamSpec("upper_girdle_length", 0.55, 0.95, 0.80,
              note="upper girdle half length as a fraction"),
    ParamSpec("lower_girdle_length", 0.60, 0.95, 0.77,
              note="lower girdle half length - controls the shape of the hearts"),
    ParamSpec("girdle_thick_main", 0.008, 0.045, 0.020,
              note="girdle thickness at the mains (thin points)"),
    ParamSpec("girdle_thick_half", 0.015, 0.070, 0.035,
              note="girdle thickness between the mains (thick points)"),
    # Pointed, and off by default.  "پایینش صفحه خورده" -- the bottom has a
    # facet on it -- is a defect report, not a request: the client's reference
    # stones come to a point, and any culet radius above zero puts a visible
    # flat plane down there.  Enable the row to let the search use one.
    ParamSpec("culet", 0.0, 0.06, 0.0, enabled=False,
              note="culet radius; 0 = pointed (what the client wants)"),
    ParamSpec("crown_twist", -3.0, 3.0, 0.0, enabled=False,
              note="crown rotation against the pavilion. Must stay 0 for hearts and arrows"),
]


@dataclass
class SearchSpace:
    """The parameters the optimiser may move, plus the fixed rest."""

    specs: list[ParamSpec] = field(default_factory=lambda: [ParamSpec(**vars(s)) for s in DEFAULT_SPACE])
    base: CutParams = field(default_factory=CutParams)

    def free(self) -> list[ParamSpec]:
        return [s for s in self.specs if not (s.fixed or self._inert(s))]

    def _inert(self, spec: ParamSpec) -> bool:
        """Parameters the base makes meaningless, whatever the sheet says.

        A straight girdle has one thickness, so searching for a second one is
        a dimension that changes nothing -- and an unused dimension is not
        free: differential evolution still has to explore it, so it costs
        generations that would otherwise go into the angles.
        """
        return spec.name == "girdle_thick_half" and self.base.flat_girdle

    def enable(self, name: str) -> "SearchSpace":
        """Let the optimiser move a row the sheet ships switched off."""
        for spec in self.specs:
            if spec.name == name:
                spec.enabled = True
        return self

    def lock(self, name: str, value: float) -> "SearchSpace":
        """Pin one parameter to a value the client chose, in place.

        The row stays in the sheet so it is still reported; it just stops
        being searched.
        """
        for spec in self.specs:
            if spec.name == name:
                spec.enabled = False
                spec.default = float(value)
        setattr(self.base, name, float(value))
        return self

    def bounds(self) -> tuple[np.ndarray, np.ndarray]:
        free = self.free()
        return (
            np.array([s.low for s in free], dtype=np.float64),
            np.array([s.high for s in free], dtype=np.float64),
        )

    def start(self) -> np.ndarray:
        return np.array([s.default for s in self.free()], dtype=np.float64)

    def to_params(self, vector: np.ndarray) -> CutParams:
        """Turn a search vector into a stone.

        Fixed parameters come from ``base``, so a spreadsheet that disables a
        row still produces a complete, valid stone.

        A spec's own default only fills in where ``base`` is silent -- i.e.
        where it still holds the dataclass default.  Otherwise passing
        ``SearchSpace(base=CutParams(n_base=16))`` would be quietly overwritten
        by the spec's 8 and the caller would get a stone they did not ask for.
        """
        values = self.base.to_dict()
        untouched = CutParams().to_dict()
        for spec in self.specs:
            if spec.fixed and values.get(spec.name) == untouched.get(spec.name):
                values[spec.name] = _cast(spec, spec.default)
        for spec, v in zip(self.free(), np.atleast_1d(vector)):
            values[spec.name] = _cast(spec, spec.clamp(v))
        return CutParams.from_dict(values)

    def to_vector(self, params: CutParams) -> np.ndarray:
        d = params.to_dict()
        return np.array([float(d[s.name]) for s in self.free()], dtype=np.float64)


def _cast(spec: ParamSpec, value: float):
    if spec.kind == "int":
        return int(round(value))
    if spec.kind == "bool":
        return bool(value >= 0.5)
    return float(value)


# --------------------------------------------------------------------------
# spreadsheet I/O
# --------------------------------------------------------------------------

_HEADER = ["parameter", "enabled", "min", "max", "start", "type", "step", "note"]


def _rows(space: SearchSpace) -> list[list]:
    return [
        [s.name, int(bool(s.enabled)), s.low, s.high, s.default, s.kind, s.step, s.note]
        for s in space.specs
    ]


def write_template(path, space: SearchSpace | None = None) -> Path:
    """Write the parameter sheet.  ``.xlsx`` if openpyxl is present, else CSV."""
    space = space or SearchSpace()
    path = Path(path)
    rows = _rows(space)

    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            path = path.with_suffix(".csv")
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "parameters"
            ws.append(_HEADER)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="333F50")
                cell.alignment = Alignment(horizontal="center")
            for row in rows:
                ws.append(row)
            widths = [22, 9, 10, 10, 10, 8, 8, 70]
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
            ws.freeze_panes = "A2"
            wb.save(path)
            return path

    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(_HEADER)
        writer.writerows(rows)
    return path


def load_space(path, base: CutParams | None = None) -> SearchSpace:
    """Read a parameter sheet back.

    Unknown parameter names are reported rather than ignored -- a typo in the
    spreadsheet silently dropping a parameter from the search would be very
    hard to notice from the results.
    """
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, data_only=True).active
        raw = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            raw = list(csv.reader(fh))

    if not raw:
        raise ValueError(f"{path} is empty")
    header = [str(h).strip().lower() for h in raw[0]]
    idx = {name: header.index(name) for name in _HEADER if name in header}
    missing = [c for c in ("parameter", "min", "max") if c not in idx]
    if missing:
        raise ValueError(f"{path} is missing column(s): {', '.join(missing)}")

    known = set(CutParams.__dataclass_fields__)
    specs: list[ParamSpec] = []
    unknown: list[str] = []
    for row in raw[1:]:
        if not row or row[idx["parameter"]] in (None, ""):
            continue
        name = str(row[idx["parameter"]]).strip()
        if name not in known:
            unknown.append(name)
            continue

        def get(col, default=None):
            if col not in idx or idx[col] >= len(row) or row[idx[col]] in (None, ""):
                return default
            return row[idx[col]]

        low = float(get("min", 0.0))
        high = float(get("max", 0.0))
        specs.append(
            ParamSpec(
                name=name,
                low=low,
                high=high,
                default=float(get("start", (low + high) / 2.0)),
                kind=str(get("type", "float")).strip().lower(),
                step=float(get("step", 0.0) or 0.0),
                enabled=str(get("enabled", 1)).strip().lower() not in ("0", "false", "no", ""),
                note=str(get("note", "") or ""),
            )
        )

    if unknown:
        raise ValueError(
            f"{path} lists parameter(s) the model does not have: {', '.join(unknown)}. "
            f"Valid names: {', '.join(sorted(known))}"
        )
    if not specs:
        raise ValueError(f"{path} contained no usable parameter rows")

    return SearchSpace(specs=specs, base=base or CutParams())
